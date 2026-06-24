import pandas as pd
import openpyxl
import configparser
import datetime
import ast
from collections import defaultdict
from helpers.utilities import get_directory, get_int, get_split_course
from helpers.utilities import get_format_headers, get_replace_header
from helpers.utilities import get_sheet_headers, get_filepath
from helpers.utilities import get_row_info, get_input


class Book:
    """Simplified version of the full Book class to store necessary information for emails."""

    def __init__(self, title, author, edition, year, access):
        self.title = title
        self.author = author
        self.edition = edition
        self.year = get_int(year)
        self.access = access


class Instructor:
    """Class to store book data related to sending emails."""

    def __init__(self, name, email, course, section, book):
        self.name = name
        self.email = email
        self.data = {course: {section: [book]}}

    def add_book(self, course, section, book):
        if course in self.data and section in self.data[course]:
            self.data[course][section].append(book)
        elif course in self.data:
            self.data[course][section] = [book]
        else:
            self.data[course] = {section: [book]}


def update_excel(directory, data, sheet_name):
    """Updates the main excel sheet with which rows have been parsed into emails."""
    full_config = configparser.ConfigParser()
    full_config.read("config.ini")
    # email_config = full_config["Email"]

    header_dict, header_list = get_sheet_headers()
    workbook = openpyxl.load_workbook(directory)
    worksheet = workbook[sheet_name]

    # all this does is just takes the index of all the rows that were processed
    # and then adds the current time / date to the date emailed cell
    for idx, row in enumerate(worksheet.iter_rows(min_row=2)):
        if idx in data:
            for cell in row:
                col = cell.column_letter
                if worksheet[f"{col}1"].value == header_dict["EmailDate"]:
                    cell.value = datetime.datetime.now()
                    break
    workbook.save(directory)


def write_to_excel(directory, export_data, sheetname):
    """Exports given data to given directory with given sheetname."""
    dataframe = pd.DataFrame(data=export_data)
    writer = pd.ExcelWriter(directory, engine="xlsxwriter")
    dataframe.to_excel(
        writer, sheet_name=sheetname, startrow=1, header=False, index=False
    )
    worksheet = writer.sheets[sheetname]
    max_row, max_col = dataframe.shape
    column_settings = []
    for header in dataframe.columns:
        column_settings.append({"header": header})
    worksheet.add_table(0, 0, max_row, max_col - 1, {"columns": column_settings})
    worksheet.set_column(0, max_col - 1, 12)
    writer.close()


# TODO
# rewrite this into smaller functions to break it apart a little
def create_email_excel(input_sheet=None, file_name=""):
    """Creates an excel sheet to be formatted for usage with PowerAutomate function."""
    full_config = configparser.ConfigParser()
    full_config.read("config.ini")
    email_config = full_config["Email"]

    text_config = configparser.ConfigParser()
    text_config.read("helpers/ini/emails.ini")
    text_config = text_config["Main"]

    if file_name == "":
        input_path = get_directory("Input", email_config)
    else:
        input_path = get_filepath(file_name)

    output_path = get_directory("Save", email_config)

    # here is the older config input
    # process_sheet = email_config["Sheetname"]

    # changed to accommodate function call from GUI
    if input_sheet is None:
        process_sheet = get_input(
            text="Input which sheet is being pulled from (i.e. Spring26): "
        )
    else:
        process_sheet = input_sheet

    data = pd.read_excel(input_path, sheet_name=process_sheet)

    instructors_list = []
    course_amt = 0
    section_list = []

    # formatting headers : course num (0), section num (1), instructor/email (2)
    # $ -> course number, & -> section number, ^ -> "instructor/email"
    headers = get_format_headers()
    header_dict, header_list = get_sheet_headers()

    for header, col in data.items():
        if "Course" in header:
            course_num = get_int(header[7])
            if course_num > course_amt:
                course_amt = course_num
                section_list.append(1)
            elif "Section" in header:
                # this might be a culprit later
                # TODO
                # update this string slice to look at other headers ?
                # i doubt the headers will be changed but you never know
                section_num = get_int(
                    "".join([char for char in header[9:] if char.isdigit()])
                )
                if section_num > section_list[course_num - 1]:
                    section_list[course_num - 1] = section_num

    book_idx = []

    for idx, row in data.iterrows():

        # this checks if it has already been emailed
        date_emailed = (
            row[header_dict["EmailDate"]]
            if not pd.isna(row[header_dict["EmailDate"]])
            else ""
        )
        if date_emailed != "":
            continue

        # this checks if it is even able to be emailed (as in has all the
        # proper data & in reading list)
        formats_list = (
            row[header_dict["ReadingList"]]
            if not pd.isna(row[header_dict["ReadingList"]])
            else ""
        )
        if formats_list is False or formats_list == "False":
            continue

        # if we havent tracked this row yet, add it to email index we are
        # making
        if idx not in book_idx:
            book_idx.append(idx)

        # gather book data first
        title = get_row_info(row, "Title")
        author = get_row_info(row, "Author")
        edition = get_row_info(row, "Edition")
        year = get_row_info(row, "Year")
        # isbn = get_row_info(row, "ISBN")

        access = {}

        # basically we grab all the MMS IDs; if they're empty, we dont need the rest of the data
        # if we have it, we grab the data we need to include!
        ebook_mms = get_row_info(row, "EbookMMSId")
        print1_mms = get_row_info(row, "PrintMMSId1")
        print2_mms = get_row_info(row, "PrintMMSId2")
        audio_mms = get_row_info(row, "AudioMMSId")

        if ebook_mms != "":
            access["Ebook"] = {
                "Number": get_row_info(row, "EbookUsers"),
                "Link": get_row_info(row, "EbookLink"),
                "CDL": get_row_info(row, "IsCDL"),
            }

        if print1_mms != "":
            access["Print1"] = {
                "Number": get_row_info(row, "PrintCopies1"),
                "Link": get_row_info(row, "PrintLink1"),
            }

        if print2_mms != "":
            access["Print2"] = {
                "Number": get_row_info(row, "PrintCopies2"),
                "Link": get_row_info(row, "PrintLink2"),
            }

        if audio_mms != "":
            access["Audio"] = {"Link": get_row_info(row, "AudioLink")}

        # parse data into book class
        book = Book(title, author, edition, year, access)

        # then iterate through course / section rows for professors
        for course in range(1, course_amt + 1):
            course_name = row[f"{get_replace_header(headers[0], course)}"]
            if pd.isna(course_name):
                continue

            for section in range(1, section_list[course - 1] + 1):
                section_name = row[f"Course {course}, Section {section}"]
                if pd.isna(section_name):
                    continue
                section_name = get_int(section_name)
                section_inst = row[f'{get_replace_header(headers[2],
                                                         course,
                                                         section,
                                                         "Instructor")}']
                section_email = row[f'{get_replace_header(headers[2],
                                                          course,
                                                          section,
                                                          "Email")}']

                if "STAFF" in section_inst:
                    continue

                if pd.isna(section_inst):
                    continue

                inst_found = False
                for instructor in instructors_list:
                    if section_inst == instructor.name:
                        instructor.add_book(course_name, section_name, book)
                        inst_found = True
                        break
                if inst_found:
                    continue
                new_inst = Instructor(
                    section_inst, section_email, course_name, section_name, book
                )
                instructors_list.append(new_inst)

    # this is the final data output that will be parsed into the Excel sheet
    final_data = {"Name": [], "Email": [], "Book Output": []}

    # the rest of this is iterating through each professor and building the HTML markdown for the email itself
    # based on the information we have, we include/exclude certain pieces of
    # written information
    for instructor in instructors_list:
        final_data["Name"].append(instructor.name)
        final_data["Email"].append(instructor.email)

        email_str = ""
        scanned_appear = False
        physical_appear = False
        ebook_appear = False

        for k, course_code in enumerate(instructor.data):
            if k != 0:
                email_str += "<br>"

            course = get_split_course(course_code)
            email_str += "<br>"
            email_str += text_config["Link"].replace("[0]", course[0]).replace("[1]", course[1])
            # email_str += f'<b>{
            #     course[0]} {
            #     course[1]}</b><br>Students can find all library materials by <a href="https://search.library.oregonstate.edu/discovery/search?query=any,contains,{
            #     course[0]}%20{
            #         course[1]}&tab=CourseReserves&search_scope=CourseReserves&vid=01ALLIANCE_OSU:OSU&lang=en&offset=0">searching course reserves for {
            #             course[0]} {
            #                 course[1]}</a>.<br><ul>'

            # TODO
            # combining class codes ?

            # book_dict = {book_str1: [section1, section2],
            #                book_str2: [section1], etc.}

            # and then to implement this on a course level, compare course section_str
            # if they are the same, then combine the course into one set
            # this will require having multiple find all materials phrases?
            # can also skip over this, the section space save is more important
            book_dict = defaultdict(list)

            for section in instructor.data[course_code]:
                for book in instructor.data[course_code][section]:
                    book_str = ""
                    book_str += f"<li><em>{book.title}</em>, {book.author}"
                    book_str += f", {
                            book.edition} Edition" if not (book.edition == "") else ""
                    book_str += (
                        f", {book.year}</li>" if not (book.year == "") else "</li>"
                    )

                    book_str += "<ul>"

                    print(book.access)
                    for access_data in book.access:
                        link = book.access[access_data]["Link"]
                        print(book.access[access_data])

                        if access_data == "Ebook":
                            user_num = book.access[access_data]["Number"]

                            # CDL
                            if book.access[access_data]["CDL"]:
                                scanned_appear = True
                                type_str = "Scanned Book"
                            # ebook
                            else:
                                ebook_appear = True
                                type_str = "Ebook"

                            if user_num == "unlimited":
                                book_str += f'<li><a href="{link}">{type_str}</a>: unlimited simultaneous users</li>'
                            else:
                                if not pd.isna(user_num) and user_num is not None:
                                    user_int = get_int(user_num)
                                    user_str = "users" if user_int != 1 else "user"
                                else:
                                    user_num = "Unknown"
                                    user_str = "users"
                                book_str += f'<li><a href="{link}">{type_str}</a>: {user_num} simultaneous {user_str}</li>'

                        elif access_data == "Print1" or access_data == "Print2":
                            physical_appear = True
                            copy_int = get_int(book.access[access_data]["Number"])
                            copy_str = "copies" if copy_int != 1 else "copy"
                            if access_data == "Print1":
                                type_str = "Print"
                            else:
                                type_str = "Print (Alt.)"
                            book_str += f'<li><a href="{link}">{type_str}</a>: {copy_int} {copy_str} in Course Reserves</li>'

                        elif access_data == "Audio":
                            book_str += f'<li><a href="{link}">Audiobook</a> (limited users)</li>'

                    # closing the unordered list for the book access
                    book_str += "</ul>"
                    book_dict[book_str].append(section)

            # iterate through this dictionary, recompile a section dict; ensure sections are sorted
            # section_dict = {[section1, section2]: book_str1, book_str5,
            #                 [section1]: book_str2, book_str3, etc.}
            section_dict = defaultdict(list)
            for book_str in book_dict:
                section_dict[f"{book_dict[book_str]}"].append(book_str)

            for sections in section_dict:
                section_str = ""
                sections = ast.literal_eval(sections)
                sec_len = len(sections)
                if sec_len == 1:
                    section_str += "<li>Section "
                    section_str += f"{sections[0]}"
                elif sec_len == 2:
                    section_str += "<li>Sections "
                    section_str += f"{sections[0]} & {sections[1]}"
                elif sec_len >= 3:
                    section_str += "<li>Sections "
                    for num, section in enumerate(sections):
                        if num == 0:
                            section_str += f"{section}"
                        elif num + 1 == sec_len:
                            section_str += f", & {section}"
                        else:
                            section_str += f", {section}"
                section_str += "<ul>"
                for book_list in section_dict[f"{sections}"]:
                    for book_str in book_list:
                        section_str += book_str
                section_str += "</ul>"
                email_str += section_str

            # closing the unordered list for the course
            email_str += "</ul>"

        # checking whether or not these specific cases
        # have appeared for this professor
        # Scanned was updated with new language according to Pre-Purchasing Email Notification Doc
        # Ebook language was also added according to same doc mentioned above
        if scanned_appear:
            email_str += "<br>"
            email_str += text_config["scanned"]
            # email_str += '<br>Scanned books are first come, first served, for one hour at a time and use a waitlist. There is no limit to the number of renewals if no one is in the waitlist. If you would like to increase the number of simultaneous users, you may provide additional print copies for us to sequester. For every print copy we sequester, we can allow one online user in accordance with <a href="https://www.controlleddigitallending.org/">controlled digital lending</a> principles.'
        if scanned_appear and physical_appear:
            email_str += "<br>"
        if physical_appear:
            email_str += "<br>"
            email_str += text_config["physical"]
            # email_str += "<br>Physical copies are available for checkout at the Borrowing & Information desk for three hours at a time."
        if (scanned_appear and ebook_appear) or (physical_appear and ebook_appear):
            email_str += "<br>"
        if ebook_appear:
            email_str += "<br>"
            email_str += text_config["ebook"]
            # email_str += "<br>When purchasing ebooks, we prioritize unlimited-user licenses, but these are not always available."

        final_data["Book Output"].append(email_str)

    write_to_excel(output_path, final_data, "Email List")
    update_excel(input_path, book_idx, process_sheet)
    print("Email sheet exported. Main sheet updated.")


# the goal here should be to load book data into a book object,
# then for each instance of an instructor for that book,
# assign the book reference to them with the course + section

# if instructor already exists, just add the other information
# otherwise, make new instructor
